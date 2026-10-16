"""Assemble the STQA mini project report (.docx) in the order of "STQA Mini Project Report Format.pdf".

    python build_report.py [output.docx]

Order (as in the college format):
  Title, Problem Definition, Prepared By, Prerequisite, Software Requirements, Hardware Requirement,
  Theory, Test Plan (Introduction, 1 Test Strategy ... 7 Test Deliverables, with 6 Bug Taxonomy),
  Test Report for Test Cases (Sr. No. / Test Objective / Steps / Expected / Actual / Status),
  Screenshot of Code and Output, GUI screenshots,
  followed by the remaining deliverables named in the title: Bug Report, Exploratory Testing,
  Test Summary Report and Conclusion.

Colour document: blue table headers, green Pass / red Fail cells, colour screenshots.
"""

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).parent
SHOTS = ROOT / "screenshots"
HEADER_FILL, KEY_FILL = "4F81BD", "DBE5F1"
STATUS_FILL = {"Pass": ("C6EFCE", RGBColor(0x00, 0x61, 0x00)), "Fail": ("FFC7CE", RGBColor(0x9C, 0x00, 0x06))}
BLACK = RGBColor(0, 0, 0)
FONT = "Times New Roman"
TEAM = [("Soniya Pawar", "[Roll No]"), ("Sakshi Jagadale", "[Roll No]"), ("Priti Merurkar", "[Roll No]")]
APP = "EduFlow – Academic Management & Productivity Portal"
TITLE = ("Create a small web-based application by selecting relevant system environments/platforms and "
         "programming languages. Narrate concise Test Plan consisting features to be tested and bug taxonomy. "
         "Narrate scripts in order to perform regression tests. Identify the bugs using Selenium WebDriver and "
         "IDE and generate test reports encompassing exploratory testing.")

doc = Document()
figure_no = 0


# ------------------------------------------------------------------ styling helpers
def set_font(style_or_run, size=None, bold=None):
    style_or_run.font.name = FONT
    style_or_run.font.color.rgb = BLACK
    if size:
        style_or_run.font.size = size
    if bold is not None:
        style_or_run.font.bold = bold
    rfonts = style_or_run.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)  # theme fonts would override Times New Roman
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)


def setup_styles():
    normal = doc.styles["Normal"]
    set_font(normal, Pt(12))
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for lvl, size in ((1, 14), (2, 12), (3, 12)):
        h = doc.styles[f"Heading {lvl}"]
        set_font(h, Pt(size), bold=True)
        h.font.italic = False
        h.paragraph_format.space_before = Pt(12 if lvl == 1 else 8)
        h.paragraph_format.space_after = Pt(4)
    for name in ("List Bullet", "List Number"):
        set_font(doc.styles[name], Pt(12))
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    add_page_number(sec)


def add_page_number(sec):
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    set_font(run, Pt(10))


def shade(cell, hex_fill=HEADER_FILL):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(par, text, bold=False, size=None):
    """Add text with **bold**, `code` and [link](url) handled; <br> becomes a line break."""
    text = text.replace("<br>", "\n")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part):
            run = par.add_run(re.match(r"\[([^\]]+)\]", part).group(1))
        else:
            run = par.add_run(part)
        if bold:
            run.bold = True
        if size:
            run.font.size = size if not run.font.name == "Consolas" else Pt(size.pt - 1)


def add_table(rows, font_size=Pt(10), widths=None, bold_first_col=False):
    header, body = rows[0], rows[1:]
    if not any(header):  # Markdown key/value table with a blank header row
        header, body, bold_first_col = body[0], body[1:], True
    if not widths:
        # every column fits its longest word; the rest of 16.5 cm is shared by text length
        cols = [[r[i] if i < len(r) else "" for r in rows] for i in range(len(header))]
        word = [max(len(w) for c in col for w in (c.replace("*", "").replace("`", "").split() or [""])) for col in cols]
        base = [0.19 * n * font_size.pt / 10 + 0.35 for n in word]
        text = [min(max(len(c) for c in col), 60) for col in cols]
        spare = max(16.5 - sum(base), 0)
        widths = [b + spare * n / sum(text) for b, n in zip(base, text)]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], h, bold=True, size=font_size)
        if rows[0] == header:
            shade(cell)
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif i == 0:
            shade(cell, KEY_FILL)
    if rows[0] == header:
        repeat_header(t.rows[0])
    for r in body:
        cells = t.add_row().cells
        for i, v in enumerate(r[: len(header)]):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], v, size=font_size, bold=bold_first_col and i == 0)
            if bold_first_col and i == 0:
                shade(cells[i], KEY_FILL)
            status = v.strip("*")
            if status in STATUS_FILL:
                fill, colour = STATUS_FILL[status]
                shade(cells[i], fill)
                for run in cells[i].paragraphs[0].runs:
                    run.bold, run.font.color.rgb = True, colour
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
    doc.add_paragraph()
    return t


def md_tables_and_text(path: Path, skip_sections=(), level_shift=0, first_heading_level=None):
    """Convert a Markdown file (headings, paragraphs, lists, tables, code) into the document.

    Level-1 headings are dropped; sections whose heading text is in skip_sections are left out.
    """
    lines = path.read_text(encoding="utf-8").splitlines()
    i, table, code, skipping, skip_level = 0, [], None, False, 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            if skipping and level <= skip_level:
                skipping = False
            if text in skip_sections:
                skipping, skip_level = True, level
        if skipping:
            i += 1
            continue
        if code is not None:
            if line.startswith("```"):
                run = doc.add_paragraph().add_run("\n".join(code))
                run.font.name, run.font.size = "Consolas", Pt(9)
                code = None
            else:
                code.append(line)
        elif line.startswith("```"):
            code = []
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                table.append(cells)
            if i + 1 >= len(lines) or not lines[i + 1].startswith("|"):
                add_table(table, bold_first_col=table[0][0] in ("Field", ""))
                table = []
        elif line.startswith("#"):
            if level > 1:
                doc.add_heading(text, min(level - 1 + level_shift, 3))
        elif re.match(r"^\s*[*-] ", line):
            add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[*-] ", "", line))
        elif re.match(r"^\s*\d+\. ", line):
            buf = [re.sub(r"^\s*\d+\. ", "", line)]
            while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]):
                i += 1
                buf.append(lines[i].strip())
            add_inline(doc.add_paragraph(style="List Number"), " ".join(buf))
        elif line.strip() and not line.startswith("---"):
            buf = [line.strip()]
            while i + 1 < len(lines) and lines[i + 1].strip() and not re.match(r"^(\s*[*-] |\s*\d+\. |#|\||```)", lines[i + 1]):
                i += 1
                buf.append(lines[i].strip())
            add_inline(doc.add_paragraph(), " ".join(buf))
        i += 1


def point(label, text="", new_page=False):
    """A top-level item of the format: '●  Label : text' as a bold heading."""
    h = doc.add_heading(level=1)
    h.paragraph_format.page_break_before = new_page
    h.add_run("●  " + label + (" :" if text else ""))
    if text:
        r = h.add_run(" " + text)
        r.bold = False
    return h


def heading(text, level=2):
    doc.add_heading(text, level)


def para(text, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.left_indent = Cm(1)
    add_inline(p, text)
    return p


def bullets(items, style="List Bullet"):
    for it in items:
        add_inline(doc.add_paragraph(style=style), it)


def roman(items):
    for n, it in enumerate(items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.space_after = Pt(1)
        add_inline(p, f"{['i', 'ii', 'iii', 'iv', 'v', 'vi', 'vii', 'viii', 'ix', 'x'][n]}) {it}")


def image(name: str, caption: str, width=16.0, max_height=21.0):
    global figure_no
    path = SHOTS / name
    if not path.exists():
        para(f"[Missing image: {name}]")
        return
    from PIL import Image

    w, h = Image.open(path).size
    width = min(width, max_height * w / h)
    doc.add_picture(str(path), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_no += 1
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"Figure {figure_no}: {caption}")
    r.italic = True
    r.font.size = Pt(10)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def md_rows(path: Path):
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("|") and not re.fullmatch(r"\|[\s:|-]+\|", line):
            rows.append([c.strip() for c in line.strip().strip("|").split("|")])
    return rows


def gantt():
    """Schedule chart drawn as a table: shaded cells mark the days each task runs."""
    plan = [
        ("Understand the application and requirements", 1, 1),
        ("Write test plan and bug taxonomy", 1, 1),
        ("Design test cases", 1, 2),
        ("Milestone: Test specification complete", 2, 2),
        ("Build automation framework and scripts", 2, 3),
        ("Record Selenium IDE flows", 3, 3),
        ("Execute tests and regression runs", 3, 4),
        ("Exploratory testing and bug reports", 4, 4),
        ("Milestone: Test execution complete", 4, 4),
        ("Final test report and documentation", 5, 5),
        ("Milestone: Test delivery", 5, 5),
    ]
    days = 5
    t = doc.add_table(rows=1, cols=3 + days)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = ["Task Name", "Start", "End"] + [f"Day {d}" for d in range(1, days + 1)]
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.text = ""
        add_inline(c.paragraphs[0], h, bold=True, size=Pt(9))
        shade(c)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for name, s, e in plan:
        cells = t.add_row().cells
        milestone = name.startswith("Milestone")
        add_inline(cells[0].paragraphs[0], name, bold=milestone, size=Pt(9))
        add_inline(cells[1].paragraphs[0], f"Day {s}", size=Pt(9))
        add_inline(cells[2].paragraphs[0], f"Day {e}", size=Pt(9))
        for d in range(1, days + 1):
            if s <= d <= e:
                if milestone:
                    cells[2 + d].paragraphs[0].add_run("◆").font.size = Pt(10)
                    cells[2 + d].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    shade(cells[2 + d])
    for row in t.rows:
        for cell, w in zip(row.cells, [6.5, 1.6, 1.6] + [1.46] * days):
            cell.width = Cm(w)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    r = p.add_run("Shaded cells show the days on which a task is carried out; ◆ marks a milestone.")
    r.italic, r.font.size = True, Pt(10)


# ------------------------------------------------------------------ document
def build():
    setup_styles()
    results = json.loads((ROOT / "reports" / "results.json").read_text(encoding="utf-8"))
    total = len(results)
    passed = sum(r["status"] == "Pass" for r in results.values())

    # ---------------------------------------------------------------- front matter (format page 1)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("STQA MINI PROJECT REPORT")
    r.bold, r.font.size = True, Pt(16)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(f"Web Testing of {APP} using Selenium WebDriver and Selenium IDE")
    r.bold, r.font.size = True, Pt(13)

    point("Title", TITLE)
    point("Problem Definition",
          f"Perform Web testing and identify the bugs using Selenium WebDriver and IDE and generate test reports "
          f"for EduFlow – Academic Management & Productivity Portal.")
    para("EduFlow is a web application used by an institution's administrators, teachers and students. Admins "
         "manage users and classes, teachers publish assignments and grade submissions, and students submit work, "
         "manage personal tasks and track their progress. Because it handles accounts, grades and uploaded files, "
         "defects in login, access control, validation or the assignment workflow directly affect its users. The "
         "application is tested as a black box through its web interface.")

    point("Prepared By")
    add_table([["Sr. No.", "Name", "Roll No.", "Class"]] + [[str(i + 1), n, rn, "BE COMP"] for i, (n, rn) in enumerate(TEAM)],
              widths=[2, 6, 4, 4])

    point("Prerequisite", "Knowledge of Core Java")
    bullets([
        "**Core Java:** object-oriented concepts (classes, objects, inheritance, exceptions) and collections. "
        "Selenium's API (WebDriver, WebElement, By, WebDriverWait) is the same object model in Java and Python, so "
        "these concepts carry over directly.",
        "**Python 3 and pytest:** the regression scripts of this project are written in Python with the pytest "
        "framework (fixtures, assertions).",
        "Basic HTML, CSS and the DOM, and CSS / XPath locators.",
        "Software testing fundamentals: test plan, test case design, equivalence partitioning, boundary value "
        "analysis, defect life cycle.",
    ])

    point("Software Requirements")
    add_table([
        ["Software", "Version / details"],
        ["Operating system", "Windows 11 (64-bit)"],
        ["Application under test", "EduFlow (PHP 8.4.5, MySQL 8.0, Bootstrap 5.3, Chart.js 4.4)"],
        ["Web server", "PHP built-in development server (or Apache via XAMPP)"],
        ["Database server", "MySQL 8.0"],
        ["Browser", "Google Chrome 154"],
        ["Programming language", "Python 3.13"],
        ["Automation tool", "Selenium WebDriver 4.27.1 (Selenium Manager for ChromeDriver)"],
        ["Record and playback tool", "Selenium IDE browser extension"],
        ["Test framework / reporting", "pytest 8.3.4, pytest-html 4.1.1"],
        ["Editor / IDE", "Visual Studio Code"],
        ["Documentation", "Microsoft Word, Markdown"],
        ["Version control", "Git, GitHub"],
    ], widths=[5.5, 11])

    point("Hardware Requirement")
    add_table([
        ["Component", "Minimum requirement"],
        ["Processor", "Intel Core i3 / AMD Ryzen 3 or better, 64-bit"],
        ["RAM", "8 GB (Chrome, MySQL and PHP run together)"],
        ["Hard disk", "2 GB free disk space"],
        ["Display", "1366 × 768 or higher resolution"],
        ["Network", "Internet access for first-time ChromeDriver download and CDN assets"],
    ], widths=[5.5, 11])

    # ---------------------------------------------------------------- theory (format pages 1-2)
    point("Theory")
    heading("Introduction of mini project")
    para("This mini project applies the software testing life cycle to a complete, working web application. "
         "EduFlow has three roles (Admin, Teacher, Student) and features such as authentication, user and class "
         "management, assignment publishing and grading, file submission, personal task management, notifications, "
         "a calendar and analytics charts. The team planned the testing, designed 36 test cases, automated them with "
         "Selenium WebDriver, recorded smoke flows in Selenium IDE, explored the application manually, and reported "
         "the defects found against defined exit criteria.")
    para("**Software testing** is the process of executing a program to find defects and to check that it meets its "
         "requirements. **Quality assurance** is the wider set of activities that make sure a process produces "
         "quality software. The testing types used here are integration testing, system testing, regression testing "
         "(re-running the automated suite after changes) and exploratory testing (simultaneous learning, test design "
         "and execution).")

    heading("Testing Tool: Selenium")
    para("Selenium is a free (open-source) automated testing suite for web applications across different browsers "
         "and platforms. It is similar to HP Quick Test Pro (QTP, now UFT), except that Selenium focuses only on "
         "automating web-based applications. Testing done using the Selenium tool is usually referred to as Selenium "
         "Testing. Selenium is not a single tool but a suite of software, each part catering to different testing "
         "needs of an organization. It has four components:")
    bullets([
        "**Selenium Integrated Development Environment (IDE)** – records user actions in the browser and plays them back.",
        "**Selenium Remote Control (RC)** – the original Selenium 1 server that injected JavaScript into the browser; "
        "now deprecated and replaced by WebDriver.",
        "**WebDriver** – a programming interface that drives the browser directly through its native automation "
        "support (ChromeDriver, GeckoDriver) using the W3C WebDriver protocol.",
        "**Selenium Grid** – runs WebDriver tests in parallel on many machines, browsers and operating systems through "
        "a hub that distributes sessions to nodes.",
    ])

    heading("Selenium IDE")
    para("Selenium IDE (Integrated Development Environment) is the simplest tool in the Selenium Suite. It is a "
         "browser add-on (for Firefox and Chrome) that creates tests very quickly through its record-and-playback "
         "functionality. This feature is similar to that of QTP. It is effortless to install and easy to learn. "
         "Because of its simplicity, Selenium IDE should only be used as a prototyping tool, not as an overall "
         "solution for developing and maintaining complex test suites. Recordings are saved as `.side` project files.")
    para("Though Selenium IDE can be used without prior knowledge of programming, the tester should be familiar with "
         "HTML, JavaScript and the DOM (Document Object Model) to use it to its full potential. Knowledge of "
         "JavaScript is required for the Selenese command \"run script\".")
    para("Selenium IDE supports auto-complete mode when creating tests. This serves two purposes: it helps the tester "
         "enter commands more quickly, and it restricts the user from entering invalid commands. In this project "
         "Selenium IDE was used to record three smoke flows (student login, quick-add task, logout) in "
         "`selenium_ide/EduFlow.side`.")

    heading("Selenium WebDriver")
    para("WebDriver controls the browser directly. Tests locate elements with `By` locators (ID, NAME, CSS_SELECTOR, "
         "XPATH, LINK_TEXT), perform actions such as click and send_keys, and use explicit waits (`WebDriverWait` with "
         "`expected_conditions`) to synchronise with dynamic pages. All 36 automated regression tests of this project "
         "use Selenium 4 WebDriver.")

    heading("Setting up the Selenium project")
    para("In Java, a Selenium project is set up with Maven: the `selenium-java` dependency in the `pom.xml` file "
         "downloads the Java bindings and all their dependencies, and the project is then imported into IntelliJ IDEA "
         "or Eclipse. This project uses the equivalent Python bindings: `pip install -r requirements.txt` installs "
         "Selenium 4.27.1, pytest and pytest-html, and Selenium Manager downloads the ChromeDriver that matches the "
         "installed Chrome automatically. The suite is run with the single command `pytest`.")

    heading("Advantages of Selenium")
    roman([
        "It is open-source software.",
        "It supports various operating environments (Windows, Linux, Mac etc.).",
        "It supports various browsers (Chrome, Firefox, Edge, Safari, Opera etc.).",
        "It supports various programming languages (Java, Python, C#, Ruby, JavaScript, PHP).",
        "It supports parallel test execution (Selenium Grid).",
        "It uses fewer hardware resources than commercial tools.",
    ])
    heading("Disadvantages of Selenium")
    roman([
        "It supports web-based applications only.",
        "No reliable vendor support; help comes from the community.",
        "No centralized maintenance of elements/objects (no built-in object repository).",
        "Setting up the environment takes effort.",
        "Difficult to use without programming knowledge.",
        "Limited support for image-based testing.",
        "New browser features may not work properly until drivers are updated.",
        "No built-in test management or reporting facility (this project uses pytest-html for reports).",
    ])

    # ---------------------------------------------------------------- test plan (format pages 3-7)
    point("Test Plan", new_page=True)
    heading("INTRODUCTION")
    para("The Test Plan is designed to prescribe the scope, approach, resources, and schedule of all testing "
         f"activities of the project {APP}.")
    para("The plan identifies the items to be tested, the features to be tested, the types of testing to be "
         "performed, the personnel responsible for testing, the resources and schedule required to complete testing, "
         "and the risks associated with the plan. The application under test runs locally at "
         "`http://127.0.0.1:8000` with fictional demo data.")

    heading("1 TEST STRATEGY")
    heading("1.1 Scope of Testing", 3)
    heading("1.1.1 Feature to be tested", 3)
    para("All the features of EduFlow listed below, which are defined in its requirements for the three roles, "
         "need to be tested.")
    features = md_rows(ROOT / "test_plan.md")
    start = next(i for i, r in enumerate(features) if r[0] == "Module")
    feat = [["Module Name", "Applicable Roles", "Description"]] + [r for r in features[start + 1:] if len(r) == 3][:14]
    add_table(feat, widths=[4, 3, 10], bold_first_col=True)

    heading("1.1.2 Feature not to be tested", 3)
    para("These features are not tested because they are outside the web user interface or not available in the "
         "test environment:")
    bullets([
        "Email delivery of password-reset links (mail is written to a log file; no mail server)",
        "Scheduled reminder job `bin/reminders.php` (command-line job, not part of the web UI)",
        "CSV report contents and CSV user import (file contents outside the UI scope)",
        "Website performance and load testing (no load-testing tool in scope)",
        "Cross-browser testing beyond Google Chrome (only Chrome on the test machine)",
        "Attendance and parent access (not implemented in the application)",
    ])

    heading("1.2 Test Type", 3)
    para("In the project EduFlow, four types of testing are conducted:")
    bullets([
        "**Integration Testing:** UI, server and database are combined and tested as a group, e.g. a student "
        "submission appears in the teacher's review queue, and publishing an assignment creates notifications.",
        "**System Testing:** conducted on the complete, integrated system to evaluate its compliance with the "
        "specified requirements, using end-to-end journeys for every role.",
        "**Regression Testing:** the automated Selenium WebDriver suite is re-run after every fix or script change "
        "to confirm that nothing else broke.",
        "**Exploratory Testing:** time-boxed sessions probing special characters, long strings, double submits, "
        "Back/Refresh, direct URLs and small screen sizes.",
    ])

    heading("1.3 Risk and Issues", 3)
    add_table([
        ["Risk", "Mitigation"],
        ["Tests change data (new users, submissions), so a second run behaves differently",
         "Reset the database to the demo seed before every full run (`migrate.php --fresh` + `seed.php`)"],
        ["Login lockout (5 failed attempts in 15 min) blocks accounts used by other tests",
         "The lockout test uses a dedicated account; the database reset clears the counters"],
        ["Timing issues with JavaScript-loaded content cause false failures",
         "Explicit waits (`WebDriverWait`) on visible elements; no fixed sleeps"],
        ["ChromeDriver / Chrome version mismatch stops the suite",
         "Selenium Manager downloads the matching driver automatically"],
        ["Team members lack experience with Selenium",
         "Shared helper module (login, waits, screenshots) and Selenium IDE recordings as examples"],
        ["The project schedule is tight",
         "Set test priority: login, access control and the assignment workflow first"],
    ], widths=[8, 8.5])

    heading("1.4 Test Logistics", 3)
    heading("1.4.1 Who will test?", 3)
    para("The project team of three members performs all testing (no outsourced testers):", indent=True)
    add_table([["Activity", "Who", "When"]] + [r for r in md_rows(ROOT / "test_plan.md") if r[0].startswith(
        ("Requirement study", "Test case design", "Selenium", "Test execution", "Exploratory testing", "Final test"))],
        widths=[7.5, 6, 3])
    heading("1.4.2 When will test occur?", 3)
    para("The testers start test execution when all of the following inputs are ready:", indent=True)
    bullets([
        "Software (EduFlow) is available for testing and running with the demo seed data",
        "Test specification (test plan and test cases) is created",
        "Test environment is built (PHP, MySQL, Chrome, Python, Selenium)",
        "Enough human resource is available for testing",
    ])

    heading("2 TEST OBJECTIVE")
    para("The test objective is to verify the functionality of EduFlow. The project focuses on testing the academic "
         "operations such as login and access control, user and class management, assignment publishing, "
         "submission and grading, and the student productivity features, to guarantee that all these operations work "
         "normally in a real institution. In detail:")
    bullets([
        "Verify that every feature in scope works as intended for each role.",
        "Verify input validation: required fields, formats and boundary values.",
        "Verify that access control prevents unauthenticated and unauthorised access.",
        "Find and document defects with clear reproduction steps, severity and priority.",
        "Provide an automated regression suite that can be re-run after changes.",
    ])

    heading("3 TEST CRITERIA")
    heading("3.1 Suspension Criteria", 3)
    para("If the team members report that 40% or more of the test cases have failed, or a blocker is found (the "
         "application does not start, or no role can log in), testing is suspended until the development team fixes "
         "the failed cases.")
    heading("3.2 Exit Criteria", 3)
    para("Specifies the criteria that denote a successful completion of a test phase:")
    bullets([
        "Run rate is mandatory to be 100% unless a clear reason is given.",
        "Pass rate is 80%; achieving the pass rate is mandatory.",
        "No Critical defect is left open.",
    ])

    heading("4 RESOURCE PLANNING")
    heading("4.1 System Resource", 3)
    add_table([
        ["No.", "Resources", "Descriptions"],
        ["1.", "Server", "PHP 8.4 built-in web server (or Apache) and a database server with MySQL 8.0"],
        ["2.", "Test tool", "Selenium WebDriver 4.27 with Python 3.13 and pytest 8.3 for automated tests; "
                            "pytest-html 4.1 to generate the test report; Selenium IDE for record and playback"],
        ["3.", "Browser", "Google Chrome 154 (headless and normal mode)"],
        ["4.", "Network", "Local host (127.0.0.1); Internet line for the first ChromeDriver download"],
        ["5.", "Computer", "Windows 11, 64-bit, 8 GB RAM, Intel Core i3 / AMD Ryzen 3 or better"],
        ["6.", "Version control", "Git and GitHub"],
    ], widths=[1.5, 3.5, 11.5])
    heading("4.2 Human Resource", 3)
    add_table([
        ["No.", "Member", "Tasks"],
        ["1.", "Soniya Pawar\n(Test Lead)", "Manage the whole project; define the test plan and bug taxonomy; build the "
                                             "automation framework; prepare the final report"],
        ["2.", "Sakshi Jagadale\n(Test Engineer)", "Design test cases; write Selenium WebDriver scripts; execute tests; "
                                                    "exploratory testing"],
        ["3.", "Priti Merurkar\n(Test Engineer)", "Design test cases; record Selenium IDE flows; exploratory testing; "
                                                   "write bug reports"],
    ], widths=[1.5, 4.5, 10.5])

    heading("5 SCHEDULE & ESTIMATION")
    heading("5.1 All project task and estimation", 3)
    est = md_rows(ROOT / "test_plan.md")
    start = next(i for i, r in enumerate(est) if r[0] == "Task")
    rows = [["Task", "Members", "Estimate effort"]]
    for r in est[start + 1:]:
        if len(r) != 3:
            break
        rows.append([r[0], r[1], f"**{r[2].strip('*')} man-hour**" if "Total" in r[0] else f"{r[2]} man-hour"])
    add_table(rows, widths=[7.5, 6, 3])
    heading("Schedule to complete these tasks", 3)
    gantt()

    heading("6 BUG TAXONOMY")
    md_tables_and_text(ROOT / "bug_taxonomy.md", level_shift=1)

    heading("7 TEST DELIVERABLES")
    para("Test deliverables are provided as below:")
    heading("7.1 Before testing phase", 3)
    bullets(["Test plan document (`test_plan.md`)", "Test cases document (`test_cases.md`)",
             "Test design specification and bug taxonomy (`bug_taxonomy.md`)"])
    heading("7.2 During the testing", 3)
    bullets(["Test tool: Selenium WebDriver scripts (`tests/`) and Selenium IDE project (`selenium_ide/EduFlow.side`)",
             "Test data: demo seed database and upload files (`tests/files/`)",
             "Test traceability: each test function is named after its test case (e.g. `test_tc22_…` ↔ TC22)",
             "Error logs and execution logs (`reports/console_output.txt`, failure screenshots)"])
    heading("7.3 After the testing cycles is over", 3)
    bullets(["Test results / reports (`reports/test_report.html`, `reports/final_report.md`)",
             "Defect report (`bug_report.md`)", "Exploratory testing notes (`exploratory_testing.md`)",
             "Installation / test procedure guidelines (`README.md`)", "This project report"])

    # ---------------------------------------------------------------- test report table (format pages 7-9)
    point("Test Report for Test Cases inclusive of Test Procedures for identified Test Scenarios", new_page=True)
    para(f"Environment: `http://127.0.0.1:8000`, Google Chrome 154 (headless), database reset to the demo seed before "
         f"the run, demo password for all accounts `Demo@1234`. {total} test cases were executed with Selenium "
         f"WebDriver: **{passed} passed, {total - passed} failed**.")
    rows = [["Sr. No.", "Test Objective", "Steps", "Expected Result", "Actual Result", "Test Status"]]
    for r in md_rows(ROOT / "test_cases.md")[1:]:
        tc, desc, steps, expected = r[0], r[1], r[2], r[3]
        res = results[tc]
        objective = "Check " + desc[0].lower() + desc[1:] if not desc.startswith(("Admin ", "Teacher ", "Student ")) else desc
        rows.append([tc, objective, steps, expected, res["actual"].replace("\n", " "), f"**{res['status']}**"])
    add_table(rows, Pt(9), widths=[1.4, 3.0, 4.2, 3.3, 3.5, 1.6])

    # ---------------------------------------------------------------- regression scripts narration
    point("Narration of Regression Test Scripts", new_page=True)
    md_tables_and_text(ROOT / "docs" / "narration.md", level_shift=1)

    heading("Selenium IDE recorded test (Login as student)")
    side = json.loads((ROOT / "selenium_ide" / "EduFlow.side").read_text(encoding="utf-8"))
    login = side["tests"][0]
    add_table([["No.", "Command", "Target", "Value"]] +
              [[str(i + 1), c["command"], c["target"], c["value"]] for i, c in enumerate(login["commands"])],
              Pt(9), widths=[1.2, 4, 7.3, 4])
    para(f"The project `{side['name']}.side` contains three recorded flows in the suite "
         f"\"{side['suites'][0]['name']}\": " + ", ".join(t["name"] for t in side["tests"]) +
         ". All three flows passed on replay (`python selenium_ide/replay_side.py`).")

    # ---------------------------------------------------------------- screenshots (format pages 9-11)
    point("Screenshot of Code and Output", new_page=True)
    image("CODE_helpers.png", "helpers.py – reusable App class with explicit waits, login/logout and screenshot helpers", 15)
    image("CODE_conftest.png", "conftest.py – fixtures for server start-up, database reset and the Chrome driver")
    image("CODE_test_auth.png", "test_01_auth.py – login, lockout, logout and access-control tests", 15)
    image("CODE_test_teacher.png", "test_03_teacher.py – assignment publishing and grading tests")
    image("OUTPUT_console.png", "Output: console output of the full regression run with the pass/fail summary", 15, 22)
    image("OUTPUT_html_report.png", "Output: pytest-html test report (reports/test_report.html)")

    point("Screenshots of GUI", new_page=True)
    image("TC03_student_dashboard.png", "Student dashboard after a successful login (TC03)")
    image("TC22_assignment_published.png", "Teacher's assignment page after Save & publish (TC22)")
    image("TC35_progress.png", "Student progress analytics with Chart.js charts (TC35)")

    # ---------------------------------------------------------------- bugs, exploratory, summary
    point("Bug Report", new_page=True)
    md_tables_and_text(ROOT / "bug_report.md", level_shift=1)
    image("TC23_required_fields.png", "BUG-01 – no required-field message under Class & subject (TC23)")
    image("EXP03_double_submit.png", "BUG-02 – duplicate tasks created by a double-click on Add")

    point("Exploratory Testing", new_page=True)
    md_tables_and_text(ROOT / "exploratory_testing.md", skip_sections=("Screenshots",), level_shift=1)
    image("EXP01_special_chars.png", "Exploratory: HTML typed into a task title is shown as text, not executed", 14)
    image("EXP_phone_dashboard.png", "Exploratory: dashboard at 390 px phone width – no horizontal scrolling", 6)

    point("Test Summary Report", new_page=True)
    md_tables_and_text(ROOT / "reports" / "final_report.md",
                       skip_sections=("Exploratory testing findings", "Conclusion"), level_shift=1)

    point("Conclusion")
    md_tables_and_text(ROOT / "docs" / "conclusion.md")

    out = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "STQA_Mini_Project_Report_Final.docx")
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build()
