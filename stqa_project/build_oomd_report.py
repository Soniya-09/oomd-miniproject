"""Build the OOMD (UML) report for EduFlow in the layout of OOMD_Report_Hotel_Management_System.docx.

    python uml_diagrams.py        # draws the straight-line UML diagrams -> screenshots/UML_*.png
    python build_oomd_report.py [output.docx]

Layout: cover page, title page (problem statement, objectives, scope, modelling approach), then one
UML diagram per section, each starting on a new page, with Description and Key Elements.
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).parent
SHOTS = ROOT / "screenshots"
FONT = "Times New Roman"
NAVY = RGBColor(0x2B, 0x3A, 0x67)
GREY = RGBColor(0x88, 0x88, 0x88)

STUDENT = "Soniya Pawar"
ROLL = "[Roll No]"
DIVISION = "[Class / Division]"
APP = "EduFlow"
APP_FULL = "EduFlow – Academic Management & Productivity Portal"

doc = Document()


# ------------------------------------------------------------------ helpers
def font(run, size=None, bold=None, italic=None, colour=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if colour is not None:
        run.font.color.rgb = colour
    return run


def para(text="", size=11.5, bold=None, italic=None, colour=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=5, page_break=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.page_break_before = Pt(before), Pt(after), page_break
    if text:
        font(p.add_run(text), size, bold, italic, colour)
    return p


def centre(text, size=12, bold=True, colour=None, before=0, after=6):
    return para(text, size, bold, colour=colour, align=WD_ALIGN_PARAGRAPH.CENTER, before=before, after=after)


def label_line(label, value):
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    font(p.add_run(label + ": "), 12, True)
    font(p.add_run(value), 12, False)


def sub(text):
    para(text, 12.5, True, before=7, after=4).paragraph_format.keep_with_next = True


def bullets(items):
    """'Label: text' items; the label part is bold, as in the template."""
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(3)
        head, _, rest = item.partition(": ")
        if rest:
            font(p.add_run(head + ": "), 11.5, True)
            font(p.add_run(rest), 11.5, False)
        else:
            font(p.add_run(item), 11.5, False)


def diagram(number, title, definition, image, description, key_elements, max_w=15.5, max_h=15.5):
    """One UML section on its own page: heading, definition, figure, Description, Key Elements."""
    para(f"{number}. {title}", 15, True, colour=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT, after=5, page_break=True)
    para(definition, 11.5, italic=True, after=7)
    w, h = Image.open(SHOTS / image).size
    width = min(max_w, max_h * w / h)
    doc.add_picture(str(SHOTS / image), width=Cm(width))
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_after = Pt(3)
    para(f"Figure {number}: {APP} – {title}", 10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=7)
    sub("Description")
    for text in description:
        para(text)
    sub("Key Elements")
    bullets(key_elements)


def footer():
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(f"OOMD Report | {STUDENT} ({ROLL}) | Page "), 9, colour=GREY)
    run = font(p.add_run(), 9, colour=GREY)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def setup():
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(11.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.15
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.54)
    sec.top_margin = sec.bottom_margin = Cm(2.12)
    footer()


# ------------------------------------------------------------------ document
def build():
    setup()

    # ---------------------------------------------------------------- cover page
    centre("OBJECT ORIENTED MODELLING", 20, before=110, after=0)
    centre("& DESIGNING", 20, after=4)
    centre("REPORT", 15, after=25)
    centre(APP.upper(), 22, colour=NAVY, before=15, after=4)
    centre("Academic Management & Productivity Portal", 14, colour=NAVY, after=35)
    label_line("Course Code", "410241")
    para(after=6)
    label_line("Submitted by", STUDENT)
    label_line("Roll No.", ROLL)
    label_line("Class / Division", DIVISION)
    para(after=6)
    label_line("Department", "Computer")
    label_line("Academic Year", "2026-27")

    # ---------------------------------------------------------------- title page
    centre("TITLE", 12, colour=GREY, after=3).paragraph_format.page_break_before = True
    centre(APP_FULL, 20, colour=NAVY, after=4)
    para("Object Oriented Modelling and Design using UML", 13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    sub("Problem Statement")
    para("Schools and colleges still manage academic work through notice boards, WhatsApp groups, emails and paper "
         "registers. Assignments are announced in one place, submitted in another and graded on paper, so students "
         "miss deadlines, teachers cannot see who has submitted, late and duplicate submissions are hard to track, "
         "and feedback reaches students late or not at all. Students also keep their own to-do lists separately, with "
         "no single view of everything that is due.")
    para("The aim of this project is to analyse and model EduFlow, a web-based Academic Management & Productivity "
         "Portal that brings these activities into one system. Admins must be able to manage users, classes, subjects "
         "and enrolment. Teachers must be able to publish assignments, review submissions, grade them, give feedback "
         "and return work for revision. Students must be able to view and submit assignments, manage personal tasks, "
         "see a calendar of deadlines, receive notifications and track their progress through charts. All data must be "
         "stored securely in a central MySQL database, and access must be controlled by role.")

    sub("Objectives")
    bullets([
        "Single platform: assignments, submissions, grades, announcements and personal tasks in one place.",
        "Accuracy: record every submission with its time, attempt number and late status, so no work is lost or "
        "counted twice.",
        "Timely feedback: notify students automatically when work is published, received, graded or returned.",
        "Productivity: give students a dashboard, calendar and natural-language quick-add for personal tasks.",
        "Security: role-based access (Admin, Teacher, Student), hashed passwords, CSRF protection, login lockout and "
        "validated file uploads.",
    ])

    sub("Scope")
    para("The system covers user and class management, assignment publishing, submission with file validation, "
         "grading and feedback, announcements, notifications and reminders, personal tasks, the calendar and progress "
         "analytics. Its users (actors) are the Admin, Teacher and Student. The Scheduler (a cron job) is a system "
         "actor that sends reminders, and the Mail Server is an external system used for password-reset emails.")

    sub("Modelling Approach")
    para("The problem is modelled with UML in the following order: Class Diagram (static structure), State Model "
         "(life cycle of a submission), Activity Diagram (assignment workflow), Use Case Diagram (user functions), "
         "Component Diagram (software modules and interfaces), Sequence Diagram (object interaction when an "
         "assignment is submitted and graded) and finally the Deployment Diagram (physical architecture on which the "
         "modelled software runs). Each diagram is described on its own page, together with its key elements.")

    # ---------------------------------------------------------------- 1 class
    diagram(1, "Class Diagram",
            "A class diagram is a static structural UML diagram that shows the classes of a system, their attributes "
            "and operations, and the relationships between them.",
            "UML_class.png",
            [
                "This diagram models the main classes of EduFlow. The abstract class User (id, fullName, email, "
                "passwordHash, role, isActive) provides login(), logout() and changePassword(). Admin, Teacher and "
                "Student inherit from User. Admin adds createUser(), deactivateUser() and enroll(); Teacher adds "
                "createAssignment(), grade() and returnForRevision(); Student adds submit() and addTask(). Every User "
                "receives Notifications (composition) and can post Announcements.",
                "A Teacher teaches Subjects and a Student is enrolled in Classes (both many-to-many), and Subjects are "
                "taught in Classes. A Teacher creates zero or more Assignments; each Assignment (title, description, "
                "dueAt, maxMarks, allowLate, status) refers to exactly one Class and one Subject, shown by the "
                "navigable arrows, and supports publish(), close() and isOverdue().",
                "An Assignment is composed of Submissions. A Student makes zero or more Submissions; each Submission "
                "(attempt, submittedAt, isLate, marks, status) is composed of one or more SubmissionFiles and zero or "
                "more Feedback entries, and supports resubmit(). A Student also owns PersonalTasks and TaskCategories, "
                "and a TaskCategory groups zero or more PersonalTasks (aggregation).",
            ],
            [
                "Classes: User, Admin, Teacher, Student, Notification, Announcement, Subject, Class, Assignment, "
                "Submission, SubmissionFile, Feedback, PersonalTask and TaskCategory, each with attributes and "
                "operations.",
                "Generalisation: Admin, Teacher and Student inherit from the abstract class User (hollow-triangle arrow).",
                "Composition (filled diamond): User 1 to * Notification; Assignment 1 to * Submission; Submission 1 to "
                "1..* SubmissionFile and 1 to * Feedback; Student 1 to * PersonalTask and TaskCategory.",
                "Associations and multiplicity: Teacher 1 to * Assignment (creates); Student 1 to * Submission "
                "(makes); Student * to * Class (enrolled in); Teacher * to * Subject (teaches); Subject * to * Class "
                "(taught); Assignment * to 1 Class and Subject.",
                "Aggregation (hollow diamond): TaskCategory 0..1 to * PersonalTask.",
                "Visibility: Attributes are private (-) and operations are public (+).",
            ], max_h=17)

    # ---------------------------------------------------------------- 2 state
    diagram(2, "State Model (State Chart Diagram)",
            "A state model describes the different states an object passes through and the transitions between them "
            "in response to events and conditions.",
            "UML_state.png",
            [
                "This diagram shows the life cycle of a student's submission for one assignment. When the assignment "
                "is published the submission starts in the Not submitted state. If the student uploads work before the "
                "deadline it moves to Submitted, where the files can still be replaced. If the deadline passes first it "
                "moves to Missing; from there a late upload is accepted only when late submission is allowed, and the "
                "submission is marked late.",
                "When the teacher opens a submission it moves to Under review and the files are locked. The teacher "
                "either grades it (Graded, with marks and feedback) or returns it for changes (Returned, with comments). "
                "A returned submission becomes Resubmitted when the student uploads again, which increases the attempt "
                "number, and it is then reviewed or graded again.",
                "A graded submission can go back to Returned if the student requests a resubmission and this is "
                "allowed. When the assignment is closed the life cycle ends in the final state.",
            ],
            [
                "Initial and final states: The filled circle starts the life cycle when the assignment is published; "
                "the circle with a ring ends it when the assignment is closed.",
                "States: Not submitted, Missing, Submitted, Under review, Graded, Returned and Resubmitted, each with "
                "an internal note (e.g. files locked, marks + feedback, attempt + 1).",
                "Events: upload work, deadline passes, teacher opens, teacher grades, teacher returns, student uploads "
                "again, student requests resubmission.",
                "Guard conditions: [before deadline], [late allowed] and [allowed] decide whether a transition may fire.",
                "Actions: / mark late and / add comments are performed during the transition.",
                "Loop: Returned → Resubmitted → Under review → Returned allows several revision cycles.",
            ], max_h=16)

    # ---------------------------------------------------------------- 3 activity
    diagram(3, "Activity Diagram",
            "An activity diagram models the workflow of a process as a sequence of actions and decisions; swimlanes "
            "show who is responsible for each action.",
            "UML_activity.png",
            [
                "This diagram shows the complete assignment process with three swimlanes: Teacher, System and Student. "
                "The teacher creates an assignment (title, class, subject, due date, marks) and optionally attaches "
                "study material. If it is not published now it is saved as a draft and published later; both branches "
                "join at a merge node. The System then makes the assignment visible to the class and sends a "
                "notification to the students.",
                "The student opens the assignment, reads the instructions and uploads work. The System first checks "
                "whether the deadline has passed and late submission is not allowed; if so the submission is rejected "
                "and the flow ends. Otherwise the System checks the file type and size. An invalid file shows an error "
                "and the student uploads a correct file again. A valid submission is saved (marked late if it is after "
                "the deadline) and the student is notified that it was received.",
                "The teacher reviews the submission. If the work is acceptable the teacher gives marks and feedback, "
                "the System notifies the student that it is graded, and the student views the grade and feedback. If "
                "not, the teacher returns it for revision with comments; the student resubmits improved work, which is "
                "saved again and comes back to the teacher for review.",
            ],
            [
                "Swimlanes: Teacher, System and Student show who performs each action.",
                "Actions: Create assignment, attach study material, save as draft, make visible, send notification, "
                "open assignment, upload work, reject submission, save submission, notify, review, give marks and "
                "feedback, resubmit improved work, view grade and feedback.",
                "Decision nodes: Publish now?, Deadline passed and late not allowed?, File type and size valid? and "
                "Work acceptable?",
                "Merge node: The draft and publish-now branches join before the assignment is made visible.",
                "Loops: An invalid file returns to Upload work, and a returned submission goes back through Save "
                "submission to Review submission.",
                "Initial and final nodes: One start node and two end nodes: rejected submission, and grade viewed by "
                "the student.",
            ], max_h=17)

    # ---------------------------------------------------------------- 4 use case
    diagram(4, "Use Case Diagram",
            "A use case diagram shows the functions (use cases) a system provides and the actors who interact with them.",
            "UML_use-case.png",
            [
                "The system boundary is the EduFlow System. The general actor User can Log in and Manage profile; "
                "Reset password extends Log in. Admin, Teacher and Student are specialisations of User (actor "
                "generalisation), so all three inherit these use cases.",
                "The Admin manages users, manages classes & subjects, enrolls students, assigns teachers, posts "
                "announcements and views reports; Import users (CSV) extends Manage users. The Teacher posts "
                "announcements, views reports, creates and publishes assignments and reviews & grades submissions; "
                "Return for revision extends Review & grade submission. The Student views assignments, submits "
                "assignments (which includes Validate file), manages personal tasks, views the calendar and views the "
                "progress dashboard.",
                "The Scheduler (cron job) is a system actor that sends reminders. Post announcement, Publish "
                "assignment, Review & grade submission and Send reminders all include Send notification, so every "
                "important event reaches the users concerned.",
            ],
            [
                "Actors: User (general), Admin, Teacher, Student and Scheduler (cron job, «system»).",
                "Actor generalisation: Admin, Teacher and Student inherit from User (hollow-triangle arrow).",
                "System boundary: The rectangle labelled EduFlow System contains all use cases.",
                "Use cases: Log in, Reset password, Manage profile, Manage users, Import users (CSV), Manage classes & "
                "subjects, Enroll students, Assign teachers, View reports, Create assignment, Publish assignment, Review "
                "& grade submission, Return for revision, Post announcement, View assignments, Submit assignment, "
                "Validate file, Manage personal tasks, View calendar, View progress dashboard, Send notification, Send "
                "reminders.",
                "Include relationships: Submit assignment includes Validate file; Publish assignment, Review & grade "
                "submission, Post announcement and Send reminders include Send notification.",
                "Extend relationships: Reset password extends Log in; Import users (CSV) extends Manage users; Return "
                "for revision extends Review & grade submission.",
            ], max_h=19)

    # ---------------------------------------------------------------- 5 component
    diagram(5, "Component Diagram",
            "A component diagram shows how a system is divided into modular components and the interfaces through "
            "which they depend on each other.",
            "UML_component.png",
            [
                "The diagram uses three kinds of relationship: dependencies (dashed «use» arrows), associations (solid "
                "lines) and assembly connectors (ball-and-socket), as shown in the legend. In the Frontend (Browser) "
                "package, the Pages (HTML + Bootstrap), the Dashboards (Chart.js charts) and the Task & Calendar "
                "scripts each have a dependency on the App Script component, which makes the API calls, handles forms "
                "and renders tables.",
                "The App Script is joined to the backend Router by an assembly connector: the App Script requires the "
                "JSON API (HTTPS) interface (socket) and the Router provides it (ball). Inside the Backend (PHP) package "
                "the Router, Security (login, role check, CSRF) and Controllers (Admin, Teacher, Student) are linked by "
                "associations. A second assembly connector joins the Controllers, which require the Service interface, "
                "to the Services package that provides it.",
                "Inside the Services package, the Assignment and Submission services depend on the Notification "
                "Service, and the Submission Service depends on File Storage; the Reminder Job (cron) also depends on "
                "the Notification Service. A third assembly connector joins the Services package, which requires the "
                "Data access interface, to the Database Access (PDO) component. Database Access is associated with the "
                "MySQL Database (PDO / TCP 3306), and File Storage is associated with the Uploaded Files folder, where "
                "it stores the files.",
            ],
            [
                "Components: Pages, Dashboards, Task & Calendar scripts, App Script, Router, Security, Controllers, User "
                "& Class Management, Assignment, Submission, Personal Task, Notification and Analytics services, File "
                "Storage, Reminder Job, Database Access (PDO), MySQL Database and Uploaded Files.",
                "Assembly connectors (ball-and-socket): JSON API (HTTPS) between App Script and Router; Service "
                "interface between Controllers and the Services package; Data access interface between the Services "
                "package and Database Access.",
                "Associations (solid lines): Router – Security – Controllers; Database Access – MySQL Database; File "
                "Storage – Uploaded Files.",
                "Dependencies (dashed «use» arrows): Pages, Dashboards and Task & Calendar scripts on App Script; "
                "Assignment Service, Submission Service and Reminder Job on Notification Service; Submission Service on "
                "File Storage.",
                "Packages: Frontend (Browser), Backend (PHP) and the nested Services package group related components.",
                "Legend: Explains the symbols for component, association, dependency and assembly connector.",
            ], max_h=15)

    # ---------------------------------------------------------------- 6 sequence
    diagram(6, "Sequence Diagram",
            "A sequence diagram shows how objects interact over time by exchanging messages in a particular order.",
            "UML_sequence.png",
            [
                "This diagram shows a student submitting an assignment and the teacher grading it. The student opens "
                "the assignment; the Browser requests it from the Controller, which loads it from MySQL only if it is "
                "addressed to that student, and shows the instructions and due date.",
                "When the student chooses a file and clicks Submit, the Browser posts the file with the CSRF token. The "
                "Controller calls SubmissionService.submit(), which checks that the assignment is open and applies the "
                "late rule. If submission is not allowed an error is returned. Otherwise FileStorage stores the file "
                "after checking its type, real content and size; an invalid file returns an error. For a valid file the "
                "submission is saved with status submitted, NotificationService saves a \"Submission received\" "
                "notification, and the Browser shows \"Submitted successfully\".",
                "The teacher then opens the review queue, which loads the submissions for the teacher's classes, and "
                "grades the work. SubmissionService updates the status to graded and saves the feedback, "
                "NotificationService notifies the student, and the student views the marks and feedback.",
            ],
            [
                "Lifelines: Student and Teacher (actors), Browser, Controller, SubmissionService, FileStorage, MySQL "
                "(database) and NotificationService.",
                "Messages: Numbered synchronous calls (solid arrows) and return messages (dashed arrows) in time order.",
                "Combined fragments: alt fragments choose between not allowed / allowed and invalid / valid file.",
                "Self-messages: SubmissionService checks the open/late rule and FileStorage checks the file.",
                "Security steps: the CSRF token is sent with the submission and assignments are loaded only for the "
                "students they are addressed to.",
            ], max_h=16)

    # ---------------------------------------------------------------- 7 deployment (last, own page)
    diagram(7, "Deployment Diagram",
            "A deployment diagram shows the physical architecture of a system: the hardware nodes, the execution "
            "environments running on them, and the software artifacts deployed on each.",
            "UML_deployment.png",
            [
                "The User Device (PC or mobile) runs a Web Browser execution environment, which holds the HTML pages and "
                "the JavaScript & CSS artifacts of the EduFlow user interface. The browser communicates with the Web "
                "Server over HTTPS, and loads the Bootstrap 5 and Chart.js 4 libraries from an external CDN, also over "
                "HTTPS.",
                "The Web Server (XAMPP on Windows or Linux) contains three execution environments. The Apache Web Server "
                "serves only the public/ folder, whose index.php is the single entry point for every request. Apache "
                "passes requests to the PHP 8 Runtime, which runs the EduFlow Application artifact and reads its "
                "settings from the .env config artifact. The Task Scheduler (cron or Windows Task Scheduler) runs "
                "reminders.php, which calls the application to send due-soon and overdue reminders. The application "
                "writes uploaded files and logs to the storage/ folder, which is outside public/ and therefore cannot "
                "be downloaded directly from the web.",
                "The EduFlow Application connects to the Database Server over TCP port 3306, where the MySQL 8 execution "
                "environment stores the eduflow database. For password-reset emails the application uses an external "
                "Mail Server over SMTP. Keeping the database on its own node lets it be secured, backed up and scaled "
                "independently of the web server.",
            ],
            [
                "Nodes (devices): User Device (PC / Mobile), Web Server (XAMPP / Linux) and Database Server, drawn as "
                "3D boxes with the «device» stereotype.",
                "Execution environments: Web Browser, Apache Web Server, PHP 8 Runtime, Task Scheduler and MySQL 8, "
                "nested inside their devices with the «execution environment» stereotype.",
                "Artifacts: HTML pages, JavaScript & CSS, index.php, EduFlow Application, .env config, reminders.php, "
                "Uploaded files, Logs, the eduflow database and the SMTP mail service.",
                "External nodes: CDN (Bootstrap 5, Chart.js 4) and Mail Server (SMTP mail service), marked «external».",
                "Communication paths: HTTPS between browser and web server and between browser and CDN; TCP 3306 "
                "between the application and MySQL; SMTP to the mail server for password reset.",
                "Security in deployment: only public/ is exposed by Apache; source code, .env and storage/ stay outside "
                "the web root.",
            ], max_h=14)

    out = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "EduFlow_OOMD_Report.docx")
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build()
