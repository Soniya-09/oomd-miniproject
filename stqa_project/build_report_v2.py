"""Assemble the EduFlow STQA report in the layout of ProcureMax_STQA_Report_Finallll.docx.

    python build_report_v2.py [output.docx]

Layout: cover page, table of contents, Project Details, Theory, 1 Test Plan ... 11 References,
with the test-case table, screenshot evidence and results taken from the real test run.
Re-uses the table/figure helpers of build_report.py.
"""

import json
import sys

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_report as br

doc = br.doc
ROOT = br.ROOT
HEADER_FILL = "D9E2F3"  # ProcureMax's light-blue table header
_shade, _add_table = br.shade, br.add_table
br.shade = lambda cell, hex_fill=HEADER_FILL: _shade(cell, HEADER_FILL if hex_fill == "4F81BD" else hex_fill)


def _table_dark_header(rows, *args, **kwargs):
    t = _add_table(rows, *args, **kwargs)
    for cell in t.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
    return t


br.add_table = _table_dark_header
APP_SHORT = "EduFlow"
APP_FULL = "EduFlow – Academic Management & Productivity Portal"
GUIDE = "[Guide Name]"
TEAM = [(n, r) for n, r in br.TEAM if n != "Priti Merurkar"]
MODULES = [  # (first TC, last TC, module name)
    (1, 14, "Authentication"), (15, 21, "User & Class Management"),
    (22, 26, "Assignments & Grading"), (27, 28, "Assignment Submission"),
    (29, 31, "Personal Tasks"), (32, 35, "Dashboard & Analytics"), (36, 36, "Profile"),
]


# ------------------------------------------------------------------ helpers
def setup():
    for name, size in (("Normal", 12), ("List Bullet", 12), ("List Number", 12)):
        br.set_font(doc.styles[name], Pt(size))
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    for lvl, size in ((1, 16), (2, 14), (3, 12)):
        h = doc.styles[f"Heading {lvl}"]
        br.set_font(h, Pt(size), bold=True)
        h.font.italic = False
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(8 if lvl == 1 else 6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.03)
    sec.top_margin = sec.bottom_margin = Cm(2.12)


def centre(text="", size=None, bold=False, space_after=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def field(par, instr):
    run = par.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, ""), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        elif tag is None and text == "":
            el = OxmlElement("w:t")
            el.text = "Right-click and choose Update Field to build the table of contents." if "TOC" in instr else ""
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)
    return run


def label(name, text):
    """'Name: text' paragraph with a bold label, as in ProcureMax's Project Details."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(name + ":").bold = True
    if text:
        br.add_inline(p, " " + text)
    return p


def sub(text):
    """Bold run-in sub-heading such as '1.1.1 Scope of Testing'."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.add_run(text).bold = True


def module_of(tc):
    n = int(tc[2:])
    return next(m for a, b, m in MODULES if a <= n <= b)


# ------------------------------------------------------------------ document
def build():
    setup()
    results = json.loads((ROOT / "reports" / "results.json").read_text(encoding="utf-8"))
    total = len(results)
    passed = sum(r["status"] == "Pass" for r in results.values())
    failed = total - passed
    secs = sum(r.get("duration_s", 0) for r in results.values())

    # ---------------------------------------------------------------- cover page
    for _ in range(3):
        centre()
    centre("SOFTWARE TESTING AND QUALITY ASSURANCE (STQA)", 14, True)
    centre("MINI PROJECT REPORT", 16, True)
    centre()
    centre("on")
    centre()
    centre(APP_FULL, 22, True)
    centre("Web Application Testing using Selenium WebDriver, Selenium IDE, Python and pytest", 13, True)
    for _ in range(3):
        centre()
    centre("Submitted by")
    for name, roll in TEAM:
        centre(f"{name}  ({roll})", bold=True, space_after=2)
    centre()
    centre("Under the guidance of")
    centre(GUIDE, bold=True)
    for _ in range(3):
        centre()
    centre("Department of Computer Engineering", bold=True)
    centre("Academic Year: 2026 – 27", bold=True)

    # ---------------------------------------------------------------- body section with footer
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.footer.is_linked_to_previous = False
    fp = body.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"{APP_SHORT} – STQA Mini Project Report  |  Page ")
    r.font.size = Pt(10)
    field(fp, "PAGE").font.size = Pt(10)

    toc = doc.add_paragraph()
    r = toc.add_run("Table of Contents")
    r.bold, r.font.size = True, Pt(16)
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- project details
    doc.add_heading("Project Details", 1)
    label("Title", "Create and test a web-based Academic Management & Productivity Portal, EduFlow, by preparing a "
                   "concise test plan with the features to be tested and a bug taxonomy, narrating automated "
                   "regression test scripts, identifying defects using Selenium WebDriver and Selenium IDE, and "
                   "generating test reports encompassing exploratory testing.")
    label("Problem Definition", "Perform web application testing for EduFlow, identify functional, validation and "
                                "usability-related issues through systematic test cases and Selenium automation, and "
                                "generate test reports from the actual test execution.")
    label("Prepared By", "")
    br.add_table([["Roll No.", "Name"]] + [[roll, name] for name, roll in TEAM], widths=[4, 8])
    label("Prerequisite", "Knowledge of Core Java / Python, web application fundamentals (HTML, CSS, DOM) and basic "
                          "software testing concepts.")
    label("Software Requirements", "")
    br.bullets([
        "Python 3.13 with pip",
        "Selenium WebDriver 4.27.1 (Selenium Manager for ChromeDriver)",
        "pytest 8.3.4 and pytest-html 4.1.1",
        "Selenium IDE browser extension",
        "Google Chrome browser",
        "PHP 8.4 and MySQL 8.0 (to run the application under test)",
        "Visual Studio Code",
        "EduFlow web application environment",
    ])
    label("Hardware Requirement", "A computer/laptop (Intel Core i3 / AMD Ryzen 3 or better, 8 GB RAM, 2 GB free "
                                  "disk) capable of running the browser, Python, PHP, MySQL and the EduFlow application.")

    # ---------------------------------------------------------------- theory
    doc.add_heading("Theory", 1)
    doc.add_heading("Introduction of Mini Project", 2)
    br.para("EduFlow is a web-based Academic Management & Productivity Portal used as the application under test for "
            "this Software Testing and Quality Assurance (STQA) mini project. It has three roles – Admin, Teacher and "
            "Student. Admins manage users and classes, teachers publish assignments and grade submissions, and students "
            "submit work, manage personal tasks, read notifications, use a calendar and track their progress through "
            "charts. The testing work focuses on verifying these workflows, validating inputs and access control, "
            "identifying defects, and documenting the results with evidence from the real test run.")
    doc.add_heading("Testing Tool: Selenium WebDriver", 2)
    br.para("Selenium is a free, open-source suite for automating web browsers. Selenium WebDriver drives the browser "
            "directly through its native automation support (ChromeDriver) using the W3C WebDriver protocol. Tests "
            "locate elements with `By` locators (ID, NAME, CSS selector, XPath), perform actions such as click and "
            "send_keys, and use explicit waits (`WebDriverWait`) to synchronise with dynamic pages. This project uses "
            "Python Selenium tests together with pytest as the test framework and pytest-html for the HTML test report. "
            "Selenium IDE, a record-and-playback browser extension, was used to record three smoke flows.")
    doc.add_heading("Testing Approach Used in EduFlow", 2)
    br.bullets([
        "Functional testing of major EduFlow workflows for all three roles.",
        "Boundary value analysis and equivalence partitioning for input validation.",
        "Automated regression testing using Selenium WebDriver and pytest.",
        "Record-and-playback smoke tests using Selenium IDE.",
        "Browser-based validation using Google Chrome.",
        "Exploratory testing of edge cases (special characters, double submits, small screens).",
        "Generation of test execution and result reports.",
    ])

    # ---------------------------------------------------------------- 1 test plan
    doc.add_heading("1 TEST PLAN", 1)
    br.para("The Test Plan defines the testing approach, resources, test environment, scope, execution strategy and "
            "expected deliverables for testing the EduFlow web application. The plan is intended to provide a "
            "structured basis for executing and documenting the functional and automated testing activities.")
    doc.add_heading("1.1 TEST STRATEGY", 2)
    sub("1.1.1 Scope of Testing")
    br.bullets([
        "Login, authentication and password recovery.",
        "Session handling and role-based access control.",
        "User and class management (Admin).",
        "Assignment publishing and grading (Teacher).",
        "Assignment submission and file-type validation (Student).",
        "Personal tasks, dashboard, notifications, calendar, progress analytics and profile.",
    ])
    sub("1.1.2 Features to be Tested")
    br.add_table([
        ["Module", "Feature", "Testing Focus"],
        ["Login", "Authentication", "Verify valid and invalid login, empty fields, case-insensitive email and lockout"],
        ["Access Control", "Session & roles", "Verify logout, direct URL access without login and 403 for wrong role"],
        ["Password Recovery", "Forgot / reset", "Verify email format validation and invalid reset links"],
        ["Users & Classes", "Admin workflow", "Verify user creation, duplicate email, password boundaries, search, deactivation"],
        ["Assignments", "Teacher workflow", "Verify create/publish, required fields, due-date rule and grading limits"],
        ["Submissions", "Student workflow", "Verify PDF upload and rejection of disallowed file types"],
        ["Tasks & Dashboard", "Productivity", "Verify quick-add, completion, title length, card navigation, notifications"],
        ["Calendar & Progress", "Views & charts", "Verify month/agenda views, chart rendering and period filter"],
    ], widths=[3.5, 3.5, 9.9], bold_first_col=True)
    sub("1.1.3 Test Types")
    br.bullets(["Functional Testing", "Integration Testing", "System Testing", "Regression Testing",
                "UI / Navigation Validation", "Exploratory Testing"])

    doc.add_heading("1.2 TESTING ENVIRONMENT", 2)
    br.add_table([
        ["Item", "Environment / Tool"],
        ["Operating Environment", "Windows 11 (64-bit)"],
        ["Application URL", "http://127.0.0.1:8000 (PHP built-in server)"],
        ["Application Stack", "PHP 8.4.5, MySQL 8.0, Bootstrap 5.3, Chart.js 4.4"],
        ["Browser", "Google Chrome 154 (headless and normal mode)"],
        ["Programming Language", "Python 3.13"],
        ["Automation Tool", "Selenium WebDriver 4.27.1, Selenium IDE"],
        ["Test Framework / Report", "pytest 8.3.4, pytest-html 4.1.1"],
        ["IDE", "Visual Studio Code"],
        ["Test Data", "Demo seed database; password Demo@1234 for all demo accounts"],
    ], widths=[5.5, 11.4], bold_first_col=True)

    doc.add_heading("1.3 RISK AND ISSUES", 2)
    risks = br.md_rows(ROOT / "test_plan.md")
    start = next(i for i, r in enumerate(risks) if r[0] == "Risk")
    br.add_table([["Risk", "Impact / Issue", "Mitigation"]] + [r for r in risks[start + 1:] if len(r) == 3][:6],
                 widths=[6, 3.5, 7.4])

    # ---------------------------------------------------------------- 2-5
    doc.add_heading("2 TEST OBJECTIVE", 1)
    br.para("The main objective of the project is to verify the functionality of the EduFlow web application and "
            "demonstrate automated regression testing using Selenium WebDriver. The project also aims to document "
            "test cases, record results, find and report defects with clear reproduction steps, and check that "
            "validation and access control protect the users' data.")

    doc.add_heading("3 TEST CRITERIA", 1)
    doc.add_heading("3.1 Suspension Criteria", 2)
    br.bullets([
        "40% or more of the test cases fail.",
        "The EduFlow application or database is unavailable for a significant part of the planned execution.",
        "A blocking defect (no role can log in) prevents execution of the majority of the planned test cases.",
    ])
    doc.add_heading("3.2 Exit Criteria", 2)
    br.bullets([
        "Run rate is 100%: all documented test cases have been executed.",
        "Pass rate is at least 80%.",
        "No Critical defect is left open.",
        "Test results, screenshots and execution evidence have been collected.",
        "The final test report is generated and reviewed.",
    ])

    doc.add_heading("4 RESOURCE PLANNING", 1)
    br.add_table([
        ["Resource", "Requirement", "Purpose"],
        ["Hardware", "Computer/Laptop (8 GB RAM)", "Development, application execution and automated testing"],
        ["Software", "Python + Selenium + Chrome", "Automation environment"],
        ["Server", "PHP 8.4 + MySQL 8.0", "Run the EduFlow application under test"],
        ["Human – Test Lead", TEAM[0][0], "Test plan, bug taxonomy, automation framework, final report"],
        ["Human – Test Engineer", TEAM[1][0],
         "Test case design, Selenium scripts, Selenium IDE recordings, exploratory testing, bug reports"],
    ], widths=[4, 5.4, 7.5], bold_first_col=True)

    doc.add_heading("5 TEST TOOLS", 1)
    br.add_table([
        ["Tool", "Purpose"],
        ["Selenium WebDriver", "Automates browser interaction and executes the regression test scripts"],
        ["Python", "Programming language used for the Selenium automation code"],
        ["pytest", "Test framework: fixtures, assertions, test discovery and execution"],
        ["pytest-html", "Generates the HTML test execution report"],
        ["Selenium IDE", "Records and replays browser flows (EduFlow.side)"],
        ["Google Chrome", "Browser used for test execution"],
    ], widths=[5, 11.9], bold_first_col=True)

    # ---------------------------------------------------------------- 6 test cases
    doc.add_heading("6 TEST CASES", 1)
    br.para(f"The project contains {total} documented test cases (TC01–TC{total:02d}), covering Authentication, User & "
            f"Class Management, Assignments & Grading, Submissions, Personal Tasks, Dashboard & Analytics and Profile. "
            f"All {total} test cases were automated with Selenium WebDriver and executed in one regression run; the "
            f"Actual Result and Status columns are taken from that run.")
    rows = [["TC ID", "Test Objective", "Test Steps", "Expected Result", "Actual Result", "Status"]]
    for r in br.md_rows(ROOT / "test_cases.md")[1:]:
        tc, desc, steps, expected = r[0], r[1], r[2], r[3]
        res = results[tc]
        rows.append([tc, f"{module_of(tc)}: {desc}", steps, expected, res["actual"].replace("\n", " "),
                     f"**{res['status']}**"])
    br.add_table(rows, Pt(9), widths=[1.3, 3.1, 4.0, 3.3, 3.7, 1.5])

    # ---------------------------------------------------------------- 7 execution
    doc.add_heading("7 TEST EXECUTION", 1)
    br.para("The automated suite is executed with pytest. The single command `pytest` starts the EduFlow server, resets "
            "the database to the demo seed, opens Chrome through Selenium WebDriver and performs every test case, "
            "saving a screenshot on failure and writing the HTML report to `reports/test_report.html`. The Selenium IDE "
            "flows are replayed with `python selenium_ide/replay_side.py`.")
    for step in [
        "Install the requirements: `pip install -r requirements.txt`.",
        "Verify the Python, PHP and MySQL installation.",
        "Open the automation project in Visual Studio Code.",
        "Execute the test command `pytest`.",
        "Observe Selenium-driven browser execution.",
        "Verify console output and final test statistics.",
        "Open the generated pytest-html report.",
    ]:
        br.add_inline(doc.add_paragraph(style="List Number"), step)

    # ---------------------------------------------------------------- 8 screenshots
    doc.add_heading("8 SCREENSHOT EVIDENCE", 1)
    br.para("This section presents the screenshot evidence collected during the project, grouped into environment "
            "setup, code implementation, application screens, test execution and results.")
    doc.add_heading("8.1 Environment Setup", 2)
    br.image("ENV_versions.png", "Python, Selenium, pytest, PHP and MySQL version output", 14)
    br.image("TC01_admin_overview.png", "EduFlow application ready (Admin overview page)")
    doc.add_heading("8.2 Code Implementation", 2)
    br.image("CODE_helpers.png", "helpers.py – reusable App class with explicit waits, login and screenshot helpers", 15)
    br.image("CODE_conftest.png", "conftest.py – fixtures for server start-up, database reset and Chrome driver")
    br.image("CODE_test_auth.png", "Selenium test code – login, lockout and access-control tests (TC04–TC11)", 15)
    br.image("CODE_test_teacher.png", "Selenium test code – create and publish assignment (TC22)")
    doc.add_heading("8.3 Application Screens", 2)
    br.image("TC08_lockout.png", "EduFlow login page – lockout after 5 failed attempts (TC08)")
    br.image("TC03_student_dashboard.png", "EduFlow student dashboard (TC03)")
    br.image("TC22_assignment_published.png", "Teacher's assignment published (TC22)")
    br.image("TC29_quick_add.png", "Quick-add personal task (TC29)")
    br.image("TC35_progress.png", "Student progress analytics dashboard (TC35)")
    doc.add_heading("8.4 Test Execution", 2)
    br.image("OUTPUT_console.png", "Console showing the full regression run and final results", 15, 22)
    br.image("OUTPUT_side_replay.png", "Selenium IDE flows replayed – 3 of 3 passed", 13)
    br.image("TC23_required_fields.png", "Failed test TC23 – no required-field message under Class & subject (BUG-01)")
    doc.add_heading("8.5 Results", 2)
    br.image("OUTPUT_html_report.png", "pytest-html report – full suite results")

    # ---------------------------------------------------------------- 9-11
    doc.add_heading("9 TEST RESULTS", 1)
    br.para(f"Across all {total} documented test cases (TC01–TC{total:02d}), {passed} passed and {failed} failed, as "
            f"recorded in the Test Cases table in Section 6. All {total} were automated using Selenium WebDriver and "
            f"pytest and executed in a single regression run from a freshly seeded database.")
    br.add_table([
        ["Metric", "Observed Result"],
        ["Total tests", str(total)],
        ["Passed", str(passed)],
        ["Failed", f"{failed} (TC23)" if failed == 1 else str(failed)],
        ["Run rate", "100%"],
        ["Pass rate", f"{passed / total:.1%}"],
        ["Execution time", f"{secs / 60:.1f} minutes"],
        ["Selenium IDE flows", "3 of 3 passed"],
    ], widths=[6, 10.9], bold_first_col=True)
    br.para(f"Verified overall result: {passed} of {total} test cases passed and {failed} failed "
            f"({passed / total:.1%} pass rate), which meets the exit criteria (100% run rate, at least 80% pass rate, "
            f"no open Critical defect). Two Minor defects were reported:")
    br.add_table([
        ["Bug ID", "Title", "Severity", "Priority", "Status"],
        ["BUG-01", "\"Class & subject\" required error not shown under the field (found by TC23)", "Minor", "P3", "Open"],
        ["BUG-02", "Double-clicking \"Add\" in quick-add creates duplicate tasks (exploratory)", "Minor", "P3", "Open"],
    ], widths=[1.8, 9.6, 1.9, 1.7, 1.9])

    doc.add_heading("10 CONCLUSION", 1)
    br.para("This STQA mini project applied systematic software testing practices to the EduFlow web application. The "
            "report documents the testing scope, environment, resources, test cases, automation approach and results. "
            f"All {total} test cases were automated with Selenium WebDriver and pytest using explicit waits, stable "
            f"locators and automatic failure screenshots; {passed} passed and one exposed a real validation defect. "
            "Exploratory testing found a second defect and confirmed that access control, input escaping, length limits "
            "and responsive layouts work correctly.")
    br.para("The automated suite runs unattended in about eight minutes from a clean database, so it can be repeated "
            "after every change as a regression check. Selenium WebDriver with pytest proved suitable for testing a "
            "modern JavaScript-driven web application, and the suite can be extended with Selenium Grid for "
            "cross-browser testing in the future.")

    doc.add_heading("11 REFERENCES / PROJECT ARTIFACTS", 1)
    br.bullets([
        "EduFlow project application",
        "EduFlow Selenium test automation code (`tests/`)",
        "pytest configuration (`pytest.ini`, `requirements.txt`)",
        "Selenium IDE project (`selenium_ide/EduFlow.side`)",
        "EduFlow test plan, test cases and bug taxonomy (`test_plan.md`, `test_cases.md`, `bug_taxonomy.md`)",
        "Bug report and exploratory testing notes (`bug_report.md`, `exploratory_testing.md`)",
        "EduFlow test-execution report (`reports/test_report.html`)",
        "STQA mini project report reference format supplied for this submission",
    ])

    # Word asks to refresh fields (TOC, page numbers) when the file is opened
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)

    out = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "EduFlow_STQA_Report.docx")
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    build()
