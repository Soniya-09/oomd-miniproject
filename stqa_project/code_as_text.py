"""Replace the four code screenshots in section 8.2 of an STQA report with copyable code text.

    python code_as_text.py [input.docx] [output.docx]

Only the picture paragraphs of Figures 3-6 are replaced (by a one-cell table holding the same
code excerpt that render_evidence.py screenshots); captions and everything else stay as they are.
"""

import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from render_evidence import code_excerpt

ROOT = Path(__file__).parent
T = ROOT / "tests"
EXCERPTS = {  # caption start -> (file, first line of excerpt, number of lines), as in render_evidence.py
    "helpers.py": (T / "helpers.py", "class App:", 62),
    "conftest.py": (T / "conftest.py", '@pytest.fixture(scope="session", autouse=True)', 30),
    "Selenium test code – login": (T / "test_01_auth.py", "def test_tc04_wrong_password", 62),
    "Selenium test code – create": (T / "test_03_teacher.py", "def test_tc22_create_and_publish_assignment", 40),
}


def code_table(doc, code: str):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.5)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F5F5F5")
    cell._tc.get_or_add_tcPr().append(shd)
    first = True
    for line in code.splitlines():
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        pf = p.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), "Consolas")
    return t


def main():
    src = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "EduFlow_STQA_Report.docx")
    out = ROOT / (sys.argv[2] if len(sys.argv) > 2 else "EduFlow_STQA_Report_CodeText.docx")
    doc = Document(src)
    paras = doc.paragraphs
    done = 0
    for i, p in enumerate(paras[:-1]):
        if not p._p.findall(".//" + qn("w:drawing")):
            continue
        caption = paras[i + 1].text
        key = next((k for k in EXCERPTS if caption.split(": ", 1)[-1].startswith(k)), None)
        if not key:
            continue
        path, start, n = EXCERPTS[key]
        t = code_table(doc, code_excerpt(path, start, n))
        p._p.addprevious(t._tbl)          # table goes where the picture was
        p._p.getparent().remove(p._p)
        done += 1
    doc.save(out)
    print(f"Replaced {done} code screenshots -> {out}")


if __name__ == "__main__":
    main()
